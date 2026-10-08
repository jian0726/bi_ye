# -*- coding: utf-8 -*-
"""
将任务书文本转换为符合学院要求的 .docx
格式要求（依 22 级正文要求文件）：
  正文标题：宋体小三号，加粗，居中
  一级标题：宋体小四号，加粗
  正文：宋体五号
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"D:\quanbudaima\bi_ye_she_ji\docs\01-选题与方案\任务书-简柚个人博客系统的设计与实现.docx"


def set_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text, size=10.5, bold=False, align=None, indent=True, space_after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p


doc = Document()
# 页面设置
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.6)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)

# 文档默认字体
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

# ===== 标题 =====
add_para(doc, "长沙南方职业学院毕业设计任务书", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=12)

# ===== 信息表 =====
info_rows = [
    ("单位名称", "人工智能学院", "专业", "软件技术"),
    ("学生姓名", "简之航", "学    号", "202420330504"),
    ("班    级", "软件技术2-2405", "指导教师", ""),
    ("工作单位", "", "职称", ""),
    ("企业导师", "", "工作单位", ""),
    ("职务/职称", "高级软件工程师", "选题类型", "A、产品设计"),
]
table = doc.add_table(rows=0, cols=4)
table.style = "Table Grid"
for r in info_rows:
    cells = table.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = ""
        p = cells[i].paragraphs[0]
        run = p.add_run(v)
        set_font(run, size=10.5)
table.add_row()
row = table.rows[-1]
row.cells[0].paragraphs[0].add_run("毕业设计题目").font.size = Pt(10.5)
set_font(row.cells[0].paragraphs[0].runs[0], size=10.5)
merged = row.cells[1].merge(row.cells[3])
run = merged.paragraphs[0].add_run("简柚个人博客系统的设计与实现")
set_font(run, size=10.5)

doc.add_paragraph()

# ===== 一 =====
add_para(doc, "一、选题依据与意义", size=12, bold=True, indent=False, space_after=4)

add_para(doc, "（一）选题依据", size=10.5, bold=True, indent=False)
add_para(doc, "随着互联网内容生态的持续演进，个人博客作为知识沉淀与个人品牌建设的重要载体，正重新受到技术从业者与内容创作者的重视。相较于依托第三方平台的写作方式，自建博客具有内容自主可控、数据永久留存、界面可深度定制等显著优势，能够满足创作者对内容主权与个性化表达的长期需求。")
add_para(doc, "当前主流的个人博客实现方案存在明显局限：基于静态站点生成器的方案部署简便，但缺乏后端数据支撑，无法实现用户体系、评论互动、访问统计等动态功能；基于成熟内容管理系统的方案功能完备，但系统架构臃肿、二次开发成本高，界面风格固化，难以实现现代简约的视觉表现。行业需要一套既能提供完整动态交互能力、又具备精致界面表现与清晰代码结构的博客系统实现方案，以满足技术学习者与内容创作者对功能完整性与使用体验的双重诉求。本课题即针对这一需求展开设计与开发。")

add_para(doc, "（二）选题意义", size=10.5, bold=True, indent=False)
add_para(doc, "本课题设计并实现一套前后端分离的个人博客系统，其意义体现在三个层面。在技术层面，系统完整覆盖现代 Java Web 开发的核心技术链路，包括 Spring Boot 3 框架应用、MyBatis-Plus 持久层开发、MySQL 数据库设计与索引优化、Redis 缓存应用、JWT 安全认证、Vue 3 前端工程化开发、Docker 容器化部署与 Jenkins 持续集成，是对软件技术专业所学知识的系统性串联与综合实践。在工程层面，课题涵盖需求分析、概要设计、详细设计、数据库设计、编码实现、系统测试、部署运维的完整软件工程流程，能够有效锻炼技术整合能力、项目规划能力与复杂工程问题解决能力。在应用层面，系统开发完成后可直接作为个人技术博客投入使用，具备真实使用价值，且代码结构规范、文档完整，可作为后续功能迭代的技术基础，契合高职教育培养高素质技术技能人才的目标。")

# ===== 二 =====
add_para(doc, "二、毕业设计任务及要求", size=12, bold=True, indent=False, space_after=4)
add_para(doc, "（一）毕业设计任务", size=10.5, bold=True, indent=False)
add_para(doc, "1.确定毕业设计选题：", indent=False)
add_para(doc, "毕业设计选题为简柚个人博客系统的设计与实现。")
add_para(doc, "2.完成毕业设计软件系统设计与实现：", indent=False)
add_para(doc, "（1）根据毕业设计选题内容完成项目需求分析，完成概要设计、详细设计和数据库设计。")
add_para(doc, "（2）完成以下功能模块开发：")

modules = [
    ("用户注册登录", "仅支持手机号与邮箱两种注册方式，登录时系统自动识别账号类型，并校验密码或验证码。登录成功后签发 JWT 令牌，实现无状态认证与角色权限控制。"),
    ("个人中心", "支持用户修改头像、昵称、个人简介与登录密码，并查看个人评论记录与收藏记录。"),
    ("文章浏览", "提供文章列表分页展示，支持按分类、标签筛选，支持关键词全文搜索，支持按年月归档浏览，并提供文章详情页的 Markdown 渲染与代码语法高亮。"),
    ("文章评论", "支持用户发表评论与嵌套回复，评论内容经敏感词过滤后按站点配置决定是否进入审核流程，管理员可在后台进行通过、拒绝、删除操作。"),
    ("点赞收藏", "支持用户对文章进行点赞与收藏，同一用户对同一文章不可重复操作，系统实时更新文章统计数值。"),
    ("分类管理", "支持管理员新增、修改、删除文章分类，设置分类排序值与描述信息。"),
    ("标签管理", "支持管理员新增、修改、删除文章标签，设置标签颜色，文章与标签为多对多关系。"),
    ("留言板", "支持注册用户与匿名访客提交留言，管理员可在后台回复留言内容。"),
    ("友链管理", "支持访客提交友链申请，管理员审核通过后在前台友链区域展示。"),
    ("全文搜索", "基于 MySQL 8.0 的 ngram 分词全文索引实现文章标题与正文的关键词检索。"),
    ("浏览量统计", "采用 Redis 累加计数与定时任务批量落库的方式统计文章浏览量，并基于 Redis 有序集合生成热门文章排行榜。"),
    ("文章管理", "支持管理员通过 Markdown 编辑器新增、编辑、删除文章，支持草稿保存、定时发布、置顶推荐、版本历史与版本回滚。"),
    ("评论审核", "支持管理员对待审核评论执行通过、拒绝、删除操作，并支持按文章、状态、关键词组合查询。"),
    ("用户管理", "支持管理员查询用户列表，对用户执行启用、禁用、重置密码操作。"),
    ("站点设置", "支持管理员配置博客名称、站点 Logo、备案号、社交链接、页脚版权等站点级参数。"),
    ("数据看板", "以图表形式展示近 30 天访问趋势、文章分类占比、热门文章排行及文章数、评论数、用户数等核心统计指标。"),
    ("文件上传", "基于 MinIO 对象存储实现文章封面与用户头像的上传、访问与删除。"),
]
for i, (name, desc) in enumerate(modules, 1):
    add_para(doc, f"{name}：{desc}")

add_para(doc, "（3）完成项目模块测试和运行。", indent=False)
add_para(doc, "（4）完成系统部署与运维配置：编写 Docker Compose 编排文件实现多服务一键启动，配置 Jenkins 流水线实现代码提交后的自动构建与发布，配置 Nginx 反向代理与 HTTPS 证书自动续期，配置日志收集与指标监控。")
add_para(doc, "3.形成毕业设计成果。", indent=False)
add_para(doc, "4.参加毕业设计答辩。", indent=False)

add_para(doc, "（二）毕业设计要求", size=10.5, bold=True, indent=False)
add_para(doc, "1.软件系统需求清晰、设计合理；", indent=False)
add_para(doc, "2.软件系统代码规范完整、各功能模块能够正常运行；", indent=False)
add_para(doc, "3.毕业设计(正文) 1 份，要求结构完整，条理清晰，文字通畅，截图清晰，图片大小一致，表述符合行业标准；", indent=False)
add_para(doc, "4.参考文献需是近 3 年出版的刊物，且至少 5 篇。", indent=False)

# ===== 三 =====
add_para(doc, "三、设计目的与成果表现形式", size=12, bold=True, indent=False, space_after=4)
add_para(doc, "（一）毕业设计目的", size=10.5, bold=True, indent=False)
add_para(doc, "本设计旨在运用 Vue 3、Spring Boot 3 等实用技术，开发一个功能完整、界面精致的个人博客系统。目的是通过前后端分离的开发方式，掌握 B/S 架构的基本原理与 RESTful API 的实际应用，综合锻炼数据库设计、缓存应用、安全认证、前端工程化开发、容器化部署与持续集成等实践技能，最终完成一个可真实投入使用的应用系统，从而提升全栈开发能力与工程实践能力。")
add_para(doc, "（二）成果表现形式", size=10.5, bold=True, indent=False)
add_para(doc, "1.毕业设计成果形式", indent=False)
add_para(doc, "表现为毕业设计(正文)文档和软件产品。")
add_para(doc, "2.毕业设计提交成果", indent=False)
add_para(doc, "文档：毕业设计(正文)；")
add_para(doc, "软件产品：毕业设计产品、数据库及源代码。", indent=False)

# ===== 四 =====
add_para(doc, "四、实施步骤与方法", size=12, bold=True, indent=False, space_after=4)
add_para(doc, "（一）实施步骤", size=10.5, bold=True, indent=False)
steps = [
    "1.根据毕业设计选题项目进行需求分析，完成项目功能模块拆解，厘清项目设计和开发的技术要求；",
    "2.完成系统架构设计及数据库设计；",
    "3.完成软件项目模块开发；",
    "4.完成项目模块测试和运行；",
    "5.完成系统部署与运维配置；",
    "6.完成毕业设计成果整理；",
    "7.参加毕业设计答辩。",
]
for s in steps:
    add_para(doc, s, indent=False)

add_para(doc, "（二）主要方法", size=10.5, bold=True, indent=False)
methods = [
    "1.资料收集法。根据该项目需求，进行项目需求的调查、收集和资料的整理，主要收集功能需求、用户画像等。",
    "2.案例研究法。通过调查相关软件项目案例资料，从而梳理设计思路和实施步骤。",
    "3.结构化方法。根据分解与抽象的原则，按照该项目中数据处理的流程，用数据流图来建立系统的功能模型。",
    "4.项目操作法。按照项目管理的步骤进行管理，如需求整理、项目设计、项目开发、项目测试、项目运行发布。",
]
for m in methods:
    add_para(doc, m, indent=False)

# ===== 五 进度表 =====
add_para(doc, "五、进度安排", size=12, bold=True, indent=False, space_after=4)
sched = [
    ("1", "毕业设计选题", "2026年10月8日至 2026年10月15日"),
    ("2", "下达毕业设计任务书", "2026年10月16日至 2026年10月20日"),
    ("3", "毕业设计过程实施", "2026年10月21日至 2027年3月15日"),
    ("4", "毕业设计定稿、查重", "2027年3月16日至 2027年5月12日"),
    ("5", "毕业设计答辩、成绩评定", "2027年5月13日至 2027年6月2日"),
    ("6", "毕业设计材料挂网和归档", "2027年6月3日至 2027年6月15日"),
]
t2 = doc.add_table(rows=0, cols=3)
t2.style = "Table Grid"
for a, b, c in sched:
    cells = t2.add_row().cells
    for i, v in enumerate((a, b, c)):
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(v)
        set_font(run, size=10.5)

doc.add_paragraph()

# ===== 意见栏 =====
t3 = doc.add_table(rows=0, cols=2)
t3.style = "Table Grid"
opinions = [
    ("指导教师意见", "毕业设计选题科学，请按毕业设计思路和要求完成毕业设计任务，并根据毕业设计进程安排按期按质提交毕业设计作品。\n\n签字："),
    ("教研室负责人\n审批意见", "同意实施\n\n签字："),
    ("教学单位负责人\n审批意见", "同意实施\n\n签字："),
]
for a, b in opinions:
    cells = t3.add_row().cells
    p1 = cells[0].paragraphs[0]
    run = p1.add_run(a)
    set_font(run, size=10.5)
    p2 = cells[1].paragraphs[0]
    run2 = p2.add_run(b)
    set_font(run2, size=10.5)

doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run("（由指导教师填写。一式三份，学生一份，指导教师自留一份，教学单位存一份）")
set_font(run, size=9)

doc.save(OUT)
print("saved:", OUT)
