"""定位正文 docx 中「留言管理」相关段落，打印段落号 / 样式 / 文本，供定点核对。"""
import sys

from docx import Document

PATH = sys.argv[1] if len(sys.argv) > 1 else "docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx"
START = int(sys.argv[2]) if len(sys.argv) > 2 else 335
END = int(sys.argv[3]) if len(sys.argv) > 3 else 355

d = Document(PATH)
paras = d.paragraphs  # 缓存一次，禁止对其调 .index()
print(f"总段落数：{len(paras)}　总表格数：{len(d.tables)}")
print("-" * 100)
for i, p in enumerate(paras, 1):
    if START <= i <= END:
        t = p.text.strip()
        print(f"[{i}] <{p.style.name}> {t[:200]}")
