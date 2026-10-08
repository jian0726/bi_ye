"""dump docx 指定区间段落的索引/样式/文本，用于定位插入点。"""
import sys
import pathlib
from docx import Document

path = sys.argv[1]
lo, hi = int(sys.argv[2]), int(sys.argv[3])
d = Document(str(pathlib.Path(path)))
for i, p in enumerate(d.paragraphs):
    if lo <= i <= hi:
        print(f'{i:4d} [{p.style.name}] {p.text[:110]}')
print(f'--- 表格数: {len(d.tables)} ---')
