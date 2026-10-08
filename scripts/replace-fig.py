"""按题注前缀替换正文 docx 中最近的上方嵌入图片。

用法：
    python scripts/replace-fig.py "图3.19" shots/_thesis/t17-admin-messages.png           # 定位校验
    python scripts/replace-fig.py "图3.19" shots/_thesis/t17-admin-messages.png --apply   # 执行替换
"""
import pathlib
import sys

from docx import Document
from docx.oxml.ns import qn

ROOT = pathlib.Path(r'D:/quanbudaima/bi_ye_she_ji')
DOCX = ROOT / 'docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx'

caption = sys.argv[1]
png = ROOT / sys.argv[2]
apply = '--apply' in sys.argv

d = Document(str(DOCX))
cap_idx = [i for i, p in enumerate(d.paragraphs) if p.text.strip().startswith(caption)]
assert len(cap_idx) == 1, f'题注 {caption} 匹配 {len(cap_idx)} 次'
cap_idx = cap_idx[0]
print(f'题注段 P{cap_idx}: {d.paragraphs[cap_idx].text.strip()}')

img_idx, rids = None, []
for i in range(cap_idx, max(cap_idx - 5, -1), -1):
    blips = d.paragraphs[i]._p.findall('.//' + qn('a:blip'))
    if blips:
        img_idx = i
        rids = [b.get(qn('r:embed')) for b in blips]
        break
assert img_idx is not None, '题注上方未找到图片段落'
print(f'图片段 P{img_idx}，rId = {rids}')

parts = d.part.related_parts
for rid in rids:
    part = parts[rid]
    print(f'  {rid}: {part.content_type}, 原大小 {len(part.blob)} 字节')

if not apply:
    print('\n（未写盘；确认无误后加 --apply 执行）')
    sys.exit(0)

assert png.exists(), f'未找到 {png}'
new_bytes = png.read_bytes()
for rid in rids:
    parts[rid]._blob = new_bytes
d.save(str(DOCX))
print(f'\n已替换：{caption} 使用 {png.name}（{len(new_bytes)} 字节）')
