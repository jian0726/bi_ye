"""把 ip2region 离线 IP 归属地变更同步到 3 份 docx。

用法：
    python scripts/sync-docx-ipregion.py           # 仅预览
    python scripts/sync-docx-ipregion.py --apply   # 写盘
"""
import pathlib
import sys
from copy import deepcopy

from docx import Document

ROOT = pathlib.Path(r'D:/quanbudaima/bi_ye_she_ji')
DOCS = ROOT / 'docs'
APPLY = '--apply' in sys.argv


def set_cell(cell, text):
    """整格替换文本，保留首段首 run 的格式。"""
    para = cell.paragraphs[0]
    if not para.runs:
        para.add_run(text)
    else:
        para.runs[0].text = text
        for r in para.runs[1:]:
            r.text = ''
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ''


def set_para(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
    else:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ''


def insert_row_after(table, row_idx, values):
    """复制 row_idx 行插到其后，并按 values 逐格写入（保留原格式）。"""
    src = table.rows[row_idx]
    new_tr = deepcopy(src._tr)
    src._tr.addnext(new_tr)
    target = table.rows[row_idx + 1]
    assert len(target.cells) == len(values), f'列数不匹配 {len(target.cells)} vs {len(values)}'
    for cell, val in zip(target.cells, values):
        set_cell(cell, val)
    return target


changed = []

# ---------- 1. 数据库设计文档.docx：province 字段说明 ----------
f = DOCS / '02-设计文档/数据库设计文档.docx'
d = Document(str(f))
hit = 0
for tb in d.tables:
    for row in tb.rows:
        for cell in row.cells:
            if '省份（IP 归属地；本机 / 内网记为' in cell.text:
                print(f'[数据库设计文档] province 单元格: {cell.text.strip()}')
                print(f'              -> 省份（由 ip2region 离线库解析；本机 / 内网与保留地址记为“本地”）')
                if APPLY:
                    set_cell(cell, '省份（由 ip2region 离线库解析；本机 / 内网与保留地址记为“本地”）')
                hit += 1
assert hit == 1, f'province 单元格命中 {hit} 次'
if APPLY:
    d.save(str(f))
changed.append(f'数据库设计文档.docx：province 字段说明（{hit} 处）')

# ---------- 2. 系统架构设计.docx：插入 ip2region 行 + 埋点异步 ----------
f = DOCS / '02-设计文档/系统架构设计.docx'
d = Document(str(f))
hit = 0
for tb in d.tables:
    for ri, row in enumerate(tb.rows):
        if row.cells[0].text.strip() == '访问埋点':
            print(f'[系统架构设计] 第 {ri} 行「访问埋点」后插入「IP 归属地解析」')
            if APPLY:
                insert_row_after(tb, ri, [
                    'IP 归属地解析',
                    'ip2region 离线 xdb：数据文件随包内置（10.6MB），启动时整体载入内存做二分查找'
                    '（微秒级、零网络请求）；回环 / 内网与保留地址记「本地」，境外 IP 记国家名，'
                    'IPv6 记「未知」；xdb 缺失时降级且不阻断启动。访问埋点与留言属地共用该组件',
                    'ip2region 3.3.7、ProvinceResolver',
                ])
            hit += 1
            break
assert hit == 1, f'「访问埋点」行命中 {hit} 次'

hit2 = 0
for tb in d.tables:
    for row in tb.rows:
        for cell in row.cells:
            if '埋点同步写' in cell.text:
                print(f'[系统架构设计] 「埋点同步写」-> 「埋点异步写」: {cell.text.strip()[:60]}')
                if APPLY:
                    for para in cell.paragraphs:
                        for r in para.runs:
                            if '同步' in r.text:
                                r.text = r.text.replace('同步', '异步')
                hit2 += 1
if APPLY:
    d.save(str(f))
changed.append(f'系统架构设计.docx：新增「IP 归属地解析」行 + 埋点异步口径（同步→异步，{hit2} 处）')

# ---------- 3. 选题方案与技术方案.docx：技术选型表插入 ip2region ----------
f = DOCS / '01-选题与方案/选题方案与技术方案.docx'
d = Document(str(f))
hit = 0
for tb in d.tables:
    for ri, row in enumerate(tb.rows):
        cells = [c.text.strip() for c in row.cells]
        if len(cells) == 4 and cells[0] == 'MinIO' and cells[2] == '对象存储':
            print(f'[选题方案] 第 {ri} 行 MinIO 后插入 ip2region 行')
            if APPLY:
                insert_row_after(tb, ri, [
                    'ip2region', '3.3.7', 'IP 归属地离线解析',
                    '数据文件随包内置、启动时载入内存二分查找，微秒级、零外部请求，'
                    '不依赖第三方接口的可用性',
                ])
            hit += 1
            break
assert hit == 1, f'MinIO 技术选型行命中 {hit} 次'
if APPLY:
    d.save(str(f))
changed.append('选题方案与技术方案.docx：技术选型表新增 ip2region 行')

print()
if APPLY:
    for c in changed:
        print('已同步:', c)
else:
    print('（未写盘；确认后加 --apply）')
