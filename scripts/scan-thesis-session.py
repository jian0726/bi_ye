"""扫描正文 docx 中与「会话 / 令牌 / 时效」相关的段落，供增量修订定位。"""
import sys
import docx

PATH = r"docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx"
KW = ["72", "JWT", "Token", "token", "令牌", "会话", "登录态", "刷新", "有效期", "撤销", "Redis"]

d = docx.Document(PATH)
print(f"总段落 {len(d.paragraphs)}")
print(f"总表格 {len(d.tables)}")
print("-" * 70)
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if any(k in t for k in KW):
        style = p.style.name if p.style else "?"
        print(f"[{i}] ({style}) {t[:180]}")
print("-" * 70)
print("=== 表格中的相关行 ===")
for ti, tb in enumerate(d.tables):
    for ri, row in enumerate(tb.rows):
        cells = [c.text.strip() for c in row.cells]
        joined = " | ".join(cells)
        if any(k in joined for k in KW):
            print(f"表{ti} 行{ri}: {joined[:200]}")
