# -*- coding: utf-8 -*-
"""毕业设计正文 docx 生成器（学院模板体例，python-docx 原生排版）。

结构：封面 → 诚信声明 → 目录(TOC 域) → 正文(1.背景/2.设计思路/3.设计内容/4.小结)
     → 参考文献 → 附录 → 查重报告占位。正文节页码从 1 开始。
"""
import pathlib
import re
import sys

sys.path.insert(0, r'D:\quanbudaima\bi_ye_she_ji\scripts')

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from md2docx import (HEI, MONO, SONG, add_inline, base_para, png_size, set_run,
                     shade)
import thesis_text as T

ROOT = pathlib.Path(r'D:\quanbudaima\bi_ye_she_ji')
THESIS_SHOTS = ROOT / 'shots/_thesis'
FIGS = ROOT / 'shots/_figs/png'
OUT = ROOT / 'docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx'

C = WD_ALIGN_PARAGRAPH.CENTER
doc = Document()

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.6)
normal = doc.styles['Normal']
normal.font.name = SONG
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), SONG)


def para(text, size=10.5, bold=False, align=None, indent=True, after=0, before=0):
    p = base_para(doc, size=size, align=align, indent=indent, before=before, after=after)
    add_inline(p, text, size, SONG, bold)
    return p


def _outline(p, lvl):
    ppr = p._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl')
    ol.set(qn('w:val'), str(lvl))
    ppr.append(ol)


def h1(text):
    p = base_para(doc, 12, indent=False, before=10, after=5)
    add_inline(p, text, 12, SONG, bold=True)
    _outline(p, 0)


def h2(text):
    p = base_para(doc, 12, indent=False, before=7, after=4)
    add_inline(p, text, 12, SONG, bold=False)
    _outline(p, 1)


def h3(text):
    p = base_para(doc, 12, indent=False, before=6, after=3)
    add_inline(p, text, 12, SONG, bold=False)
    _outline(p, 2)


def fig(path, caption, max_h=19.0):
    w_px, h_px = png_size(path)
    w_cm = 15.2
    h_cm = w_cm * h_px / w_px
    if h_cm > max_h:
        h_cm = max_h
        w_cm = h_cm * w_px / h_px
    p = base_para(doc, align=C, indent=False, before=6)
    p.add_run().add_picture(str(path), width=Cm(w_cm))
    cap = base_para(doc, align=C, indent=False, after=8)
    add_inline(cap, caption, 10.5, HEI)


def code_block(code):
    for ln in code.strip('\n').splitlines():
        p = base_para(doc, indent=False, after=0, line=1.15)
        r = p.add_run(ln if ln else ' ')
        set_run(r, MONO, 9)
        shade(p)
    base_para(doc, indent=False, after=4, line=1.0)


def table_caption(text):
    cap = base_para(doc, align=C, indent=False, before=6, after=2)
    add_inline(cap, text, 10.5, HEI)


def md_table(rows):
    def cells(row):
        return [c.strip() for c in row.strip().strip('|').split('|')]
    parsed = [cells(r) for r in rows]
    parsed = [r for k, r in enumerate(parsed) if k != 1]
    cols = max(len(r) for r in parsed)
    t = doc.add_table(rows=0, cols=cols)
    t.style = 'Table Grid'
    for ri, row in enumerate(parsed):
        cs = t.add_row().cells
        for ci in range(cols):
            txt = row[ci] if ci < len(row) else ''
            p = cs[ci].paragraphs[0]
            p.alignment = C
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            add_inline(p, txt, 9, SONG, bold=(ri == 0))
    base_para(doc, indent=False, after=4, line=1.0)


# ============ 封面 ============
para('', after=6)
para('长沙南方职业学院', size=26, bold=True, align=C, indent=False, after=6)
para('毕  业  设  计', size=22, bold=True, align=C, indent=False, after=30)
info = [
    ('毕业设计选题类型', T.COVER['type']),
    ('毕 业 设 计 题 目', T.TITLE),
    ('指   导   教   师', ''),
    ('学   生   姓   名', '简之航'),
    ('学   生   学   号', '202420330504'),
    ('专             业', '软件技术'),
    ('班             级', '软件技术2-2405'),
]
for k, v in info:
    p = doc.add_paragraph()
    p.alignment = C
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.5
    add_inline(p, f'{k}　　{v}', 14, SONG, bold=False)
para('', after=20)
para(T.COVER['school'], size=16, bold=True, align=C, indent=False, after=4)
para(T.COVER['date'], size=14, align=C, indent=False)
doc.add_page_break()

# ============ 诚信声明 ============
para('毕业设计诚信声明', size=15, bold=True, align=C, indent=False, after=14)
para(T.DECLARATION)
para('', after=26)
para('签名：', indent=False, align=None)
para('')
para(T.COVER['date'], align=None, indent=False)
doc.add_page_break()

# ============ 目录 ============
para('目  录', size=15, bold=True, align=C, indent=False, after=12)
p = doc.add_paragraph()
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
r = OxmlElement('w:r')
t = OxmlElement('w:t')
t.text = '（在 Word 中右键此处选择“更新域”即可生成目录）'
r.append(t)
fld.append(r)
p._p.append(fld)
doc.add_page_break()

# ============ 正文节（页码从 1） ============
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.page_width, sec2.page_height = Cm(21), Cm(29.7)
sec2.top_margin = sec2.bottom_margin = Cm(2.5)
sec2.left_margin, sec2.right_margin = Cm(3.0), Cm(2.6)
pg = OxmlElement('w:pgNumType')
pg.set(qn('w:start'), '1')
sec2._sectPr.append(pg)
footer = sec2.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.alignment = C
run_a = fp.add_run()
run_b = fp.add_run()
run_c = fp.add_run()
fa = OxmlElement('w:fldChar')
fa.set(qn('w:fldCharType'), 'begin')
fb = OxmlElement('w:instrText')
fb.set(qn('xml:space'), 'preserve')
fb.text = ' PAGE '
fc = OxmlElement('w:fldChar')
fc.set(qn('w:fldCharType'), 'end')
run_a._element.append(fa)
run_b._element.append(fb)
run_c._element.append(fc)
for rr in (run_a, run_b, run_c):
    set_run(rr, size=9)

para(T.TITLE, size=15, bold=True, align=C, indent=False, after=12)

# ---- 1.背景 ----
h1('1.背景')
h2('1.1 开发背景')
for x in T.S11:
    para(x)
h2('1.2 开发目的和意义')
for x in T.S12:
    para(x)

# ---- 2.设计思路 ----
h1('2.设计思路')
h2('2.1 业务流程介绍')
for x in T.S21:
    para(x)
h2('2.2 技术方案')
for x in T.S22:
    para(x)
h2('2.3 工具设备要求')
for x in T.S23:
    para(x, indent=False)

# ---- 3.设计内容 ----
h1('3.设计内容')
h2('3.1 系统详细设计')
h3('3.1.1 系统概要说明')
for x in T.S311:
    para(x)
fig(FIGS / 'draw-module.png', '图3.1 系统功能模块图', max_h=17.0)
h3('3.1.2 系统流程图')
for x in T.S312:
    para(x)
fig(FIGS / 'draw-flow.png', '图3.2 系统流程图', max_h=17.0)

h3('3.1.3 系统架构设计')
for x in T.S313:
    para(x)

h2('3.2 数据库设计')
h3('3.2.1 数据库分析')
for x in T.S321:
    para(x)
fig(FIGS / 'db-01.png', '图3.3 系统数据库 ER 图（上：核心业务域；下：支撑域）', max_h=19.5)

h3('3.2.2 数据库表设计')
para(T.TABLE_INTRO)


def extract_tables():
    md = (ROOT / 'docs/02-设计文档/数据库设计文档.md').read_text(encoding='utf-8')
    lines = md.splitlines()
    tables = {}
    i = 0
    cur = None
    while i < len(lines):
        m = re.match(r'^### 3\.(\d+) (\w+) — (.+)$', lines[i].strip())
        if m:
            cur = {'name': m.group(2), 'cname': m.group(3), 'rows': None,
                   'index': '', 'notes': []}
            tables[int(m.group(1))] = cur
            i += 1
            continue
        if cur is not None:
            s = lines[i].strip()
            if s.startswith('| 字段名'):
                rows = [s]
                j = i + 1
                while j < len(lines) and lines[j].strip().startswith('|'):
                    rows.append(lines[j].strip())
                    j += 1
                cur['rows'] = rows
                i = j
                continue
            if s.startswith('**索引**'):
                cur['index'] = s.replace('**索引**：', '').replace('`', '')
                i += 1
                continue
            if s.startswith('**设计要点') or s.startswith('**设计说明'):
                j = i + 1
                while j < len(lines) and lines[j].strip().startswith('-'):
                    cur['notes'].append(lines[j].strip().lstrip('-').strip())
                    j += 1
                i = j
                continue
        i += 1
    return tables


tables = extract_tables()
for no in range(1, 15):
    tb = tables.get(no)
    if not tb:
        continue
    para('（%d）%s 数据表（%s）：' % (no, tb['name'], tb['cname']), indent=False)
    table_caption('表3.2.%d %s 数据表' % (no, tb['name']))
    if tb['rows']:
        md_table(tb['rows'])
    if tb['index']:
        para('索引：%s。' % tb['index'], indent=False)
    if tb['notes']:
        para('设计说明：')
        for n in tb['notes']:
            p = base_para(doc, indent=False, after=2)
            p.paragraph_format.left_indent = Pt(14)
            add_inline(p, '• ' + n)
para('除上述 14 张业务表外，设计阶段还规划了敏感词表（sensitive_word）、文章版本'
     '历史表（article_history）与后台操作日志表（operation_log）三张预留扩展表，'
     '其 DDL 设计保留在数据库设计文档中；本期系统以人工审核、访问埋点等方式承担'
     '了对应职责，故未建表。', )

# ---- 3.3 系统功能实现 ----
h2('3.3 系统功能实现')
fig_no = 3
for k, sec33 in enumerate(T.SEC33, 1):
    h3('3.3.%d %s' % (k, sec33['title']))
    for x in sec33['paras']:
        para(x)
    for png_name, caption in sec33['figs']:
        fig_no += 1
        fig(THESIS_SHOTS / (png_name + '.png'), caption.replace(
            '图3.%d' % fig_no, '').strip() if False else '图3.%d %s' % (fig_no, caption),
            max_h=18.0)
    if sec33['code']:
        code_title, code_src = sec33['code']
        para('关键代码实现：', indent=False)
        code_block(code_src)

# ---- 3.4 系统测试 ----
h2('3.4 系统测试')
for x in T.S34:
    para(x)

# ---- 4.小结 ----
h1('4.小结')
h2('4.1 毕业设计成果特点')
for x in T.S41:
    para(x)
h2('4.2 设计成果的实用价值或应用前景')
for x in T.S42:
    para(x)
h2('4.3 不足之处或遗留未予解决的问题')
for x in T.S43:
    para(x)

# ---- 参考文献 ----
h1('参考文献')
for ref in T.REFS:
    para(ref, indent=False, after=2)

# ---- 附录 ----
doc.add_page_break()
h1('附录  部分重要源码')
for title, src in T.APPENDIX:
    para(title, bold=True, indent=False, before=6, after=2)
    code_block(src)

# ---- 查重报告 ----
doc.add_page_break()
h1('查重报告')
para('（查重报告待定稿检测后粘贴于此。）')

try:
    doc.save(OUT)
    print('saved:', OUT)
except PermissionError:
    _alt = OUT.with_name(OUT.stem + '-新版' + OUT.suffix)
    doc.save(_alt)
    print('LOCKED, saved to:', _alt)
