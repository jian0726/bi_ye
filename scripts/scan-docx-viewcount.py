"""扫描 docs 下各 docx 中与浏览量相关的段落/表格单元格，供增量修订定位。"""
import pathlib
from docx import Document

BASE = pathlib.Path(r'D:/quanbudaima/bi_ye_she_ji/docs')
KEYS = ['浏览量', 'view_count', '14 张', '张表', '去重', '浏览量计数']

for p in sorted(BASE.rglob('*.docx')):
    if p.name.startswith('~$'):
        continue
    try:
        d = Document(str(p))
    except Exception as e:
        print(f'!! 打不开 {p.name}: {e}')
        continue
    hits = []
    for i, para in enumerate(d.paragraphs):
        t = para.text.strip()
        if any(k in t for k in KEYS):
            hits.append(('P', i, t[:150]))
    for ti, table in enumerate(d.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                t = cell.text.strip()
                if any(k in t for k in KEYS):
                    hits.append(('T', f'{ti}/{ri}/{ci}', t[:150]))
    if hits:
        print(f'\n===== {p.relative_to(BASE)}  (段落 {len(d.paragraphs)} / 表格 {len(d.tables)}) =====')
        for kind, loc, text in hits:
            print(f'  [{kind} {loc}] {text}')
