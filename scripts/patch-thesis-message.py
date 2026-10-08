"""正文 docx 定点修订：留言管理小节移除「回复」相关描述。

规程（踩过坑）：
- 段落定位必须排除目录区（toc 样式）；
- docx.paragraphs 每次访问都会重建列表，禁止对其调 .index()，先缓存；
- 改写只动 runs[0]，其余 run 置空，段落级格式（样式 / 缩进）原样保留。
"""
import sys

from docx import Document

PATH = sys.argv[1] if len(sys.argv) > 1 else (
    "docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx"
)

REPLACEMENTS = [
    (
        "留言管理列出全站留言",
        "留言管理列出全站留言，支持按 IP 属地模糊筛选，按留言时间倒序分页展示，"
        "列表呈现留言人、内容、IP 属地与时间。后台不提供回复能力——留言板定位为一次性留言，"
        "读者留一句、博主看过即可，若再补上回复链路，留言页便成了对话区，与站点记录式的基调不符，"
        "故该能力整体移除；对违规留言可直接删除。页面效果如图3.19所示。",
    ),
    (
        "留言与前台文字雨同源",
        "留言与前台文字雨同源：前台飘落的字幅拉取的正是本表数据，"
        "因此删除留言会同步移除其在前台字幅流中的展示，不会留下悬空的引用记录。"
        "留言无需登录即可提交，游客留言的 user_id 记为 NULL、昵称缺省为“访客”，"
        "同时记录 IP 属地，供后台按属地核查。",
    ),
]

d = Document(PATH)
cache = d.paragraphs

hits = []
for p in cache:
    if p.style.name.startswith("toc"):
        continue
    text = p.text.strip()
    for prefix, new_text in REPLACEMENTS:
        if text.startswith(prefix):
            runs = p.runs
            if not runs:
                p.add_run(new_text)
            else:
                runs[0].text = new_text
                for r in runs[1:]:
                    r.text = ""
            hits.append(prefix)
            break

print("命中并改写：", hits)
if len(hits) != len(REPLACEMENTS):
    print("警告：预期改写数量与实际不符，未保存")
    sys.exit(1)

d.save(PATH)
print("已保存：", PATH)

# 重新打开复查（只看内存对象会漏掉未写入的内容）
d2 = Document(PATH)
print("复查总段落数：", len(d2.paragraphs), "　表格数：", len(d2.tables))
for i, p in enumerate(d2.paragraphs, 1):
    t = p.text.strip()
    if t.startswith("留言管理列出全站留言") or t.startswith("留言与前台文字雨同源"):
        print(f"  [{i}] {t[:120]}")
