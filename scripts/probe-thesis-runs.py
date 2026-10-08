"""探测正文 docx 指定段落的完整文本与 run 结构。"""
import pathlib
from docx import Document

P = pathlib.Path(r'D:/quanbudaima/bi_ye_she_ji/docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx')
d = Document(str(P))
for i in (88, 89, 93, 119, 124, 128, 153, 216, 219, 220, 221, 222, 330, 402):
    p = d.paragraphs[i]
    print(f'--- P{i} runs={len(p.runs)} style={p.style.name}')
    print('   TEXT:', p.text)
    for j, r in enumerate(p.runs):
        print(f'     run{j}: {r.text!r}')
print('=== 表3.2.4 article 表 view_count 行 ===')
t = d.tables[3]
for ri, row in enumerate(t.rows):
    cells = [c.text.strip() for c in row.cells]
    if 'view_count' in cells[0] or ri == 0:
        print(ri, cells)
