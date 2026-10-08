"""
正文 docx 热门文章描述清理（3 处）

段 218：5 个 run 的宋体/Consolas 混排，逐 run 替换以保住代码字体
段 223：单 run 整段替换
段 77 ：替换「Redis 缓存应用」→「Redis 键值存储应用」
"""
from docx import Document

P = r"D:\quanbudaima\bi_ye_she_ji\docs\03-正文\毕业设计正文-简柚个人博客系统的设计与实现.docx"

d = Document(P)

# ---- 段 218 ----
p = d.paragraphs[218]
assert len(p.runs) == 5, f"段218 run 数变了: {len(p.runs)}"
parts = [
    "• ",
    "idx_view",
    " 索引覆盖该字段。原设计用于支撑热门文章排行（按 ",
    "view_count",
    " 倒序取 TOP N），该功能已于设计后期整体移除，索引暂留备用、不参与当前任何查询。",
]
for r, t in zip(p.runs, parts):
    r.text = t
print("段218 已改")

# ---- 段 223 ----
p = d.paragraphs[223]
assert len(p.runs) == 1, f"段223 run 数变了: {len(p.runs)}"
p.runs[0].text = (
    "首页数据通过一次聚合接口与两次列表接口获取：列表按「置顶优先 + 发布时间倒序」分页查询，"
    "不再引入 Redis 查询缓存；“最近落笔”与时间线列表同样走分页查询。"
    "页面滚动使用 IntersectionObserver 实现元素淡入上浮，"
    "数据异步渲染完成后需要重新触发观察，避免新元素一直隐藏。"
)
print("段223 已改")

# ---- 段 77 ----
p = d.paragraphs[77]
assert len(p.runs) == 1, f"段77 run 数变了: {len(p.runs)}"
before = p.runs[0].text
if "Redis 缓存应用" in before:
    p.runs[0].text = before.replace("Redis 缓存应用", "Redis 键值存储应用")
    print("段77 已改")
else:
    print("段77 未找到目标串，跳过")

d.save(P)
print("已保存")
