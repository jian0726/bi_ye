"""正文 docx 增量修补：浏览量按独立访客去重 + 新增 t_view_record 表（14 → 15 张表）。

按段落索引定位（替换全部在插入之前完成），逐 run 赋新文本以保住代码字体（Consolas）。
"""
import copy
import pathlib

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

P = pathlib.Path(r'D:/quanbudaima/bi_ye_she_ji/docs/03-正文/毕业设计正文-简柚个人博客系统的设计与实现.docx')
d = Document(str(P))


def set_para_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)


def set_cell_text(cell, text):
    set_para_text(cell.paragraphs[0], text)


def clone_after(anchor, ref_p, text):
    new_el = copy.deepcopy(ref_p._p)
    anchor.addnext(new_el)
    p = Paragraph(new_el, ref_p._parent)
    set_para_text(p, text)
    return p._p


# ============ 1. 段落文本替换（先做，避免索引位移） ============
# P89 数据表数量
set_para_text(d.paragraphs[89],
              '在数据层面，上述业务流程由 15 张数据表支撑：用户注册登录写入用户表，文章与分类、标签通过关联表组织，'
              '点赞、收藏与浏览去重均以记录表形式与账号或访客绑定，留言、相册、音乐各自独立成表，访问日志由拦截器异步埋点。'
              '前台每次读取都基于这套数据实时组装，不需要人工维护静态页面；后台的全部管理操作最终也都落在这套表结构上，'
              '数据流转清晰、职责单一。')

# P93 数据一致性结尾句
set_para_text(d.paragraphs[93],
              '在数据一致性方面，系统对点赞、收藏这两类“同一用户对同一目标只能操作一次”的行为，'
              '采用数据库唯一索引作为防重的最终保障：t_like_record 表以 uk_user_article（user_id, target_id）'
              '与 t_collect 表的 uk_user_article（user_id, article_id）唯一索引约束，并发请求即使同时通过存在性检查，'
              '也只有一个能够插入成功，应用层捕获唯一键冲突异常后向用户返回友好提示。'
              '浏览量同属“不能重复累加”的场景，因此采用同样的思路——以浏览记录表的唯一索引完成'
              '“同一访客对同一篇文章只计一次”的去重，再由 SQL 原子自增写回计数，避免引入缓存带来的同步与降级问题。')

# P119 表数量与域划分
set_para_text(d.paragraphs[119],
              '数据库是系统全部业务数据的存储中心。简柚个人博客系统使用 MySQL 8.0，数据库名 jianyou_blog，字符集 utf8mb4。'
              '根据业务实体与它们之间的关系，共设计 15 张数据表，覆盖用户、内容、互动、资源与配置五个域：'
              '用户域包含用户表；内容域包含文章、分类、标签与文章标签关联表；'
              '互动域包含点赞记录、收藏、浏览记录与留言表；资源域包含上传资源、相册照片与背景音乐表；'
              '配置域包含友链、站点配置与访问日志表。')

# P124 冗余计数说明
set_para_text(d.paragraphs[124],
              '表结构总体上满足第三范式，但统计字段是一个重要例外：文章表直接保留 view_count、like_count、collect_count '
              '三个计数列，使列表页一次查询即可拿到展示所需的全部数据，不必对互动记录表做聚合统计。'
              '这一取舍以少量写放大换取明显的读性能提升，符合博客站点“读多写少”的实际访问特征；'
              '其中 view_count 的更新需先经浏览记录表完成独立访客去重（详见 3.2.15），'
              '再与其余计数一样由服务层在对应业务动作中同步写入，并对取消点赞做了不为负的兜底，确保冗余值与明细记录始终一致。'
              '与之相对，正文的渲染结果则刻意不做冗余：content 字段只保存 Markdown 原文，渲染交由前端完成，'
              '既减少一次写放大，也让渲染规则可以在不重写数据的前提下随时调整。')

# P128 表设计引言
set_para_text(d.paragraphs[128], '根据上述实体关系，共设计 15 张数据表。以下给出各表的字段结构、索引与设计说明。')

# P219 浏览量实现（5 runs，逐 run 替换保代码字体）
p219 = d.paragraphs[219]
assert len(p219.runs) == 5, f'P219 run 数变了: {len(p219.runs)}'
for r, t in zip(p219.runs, [
    '• 浏览量口径为「读过这篇文章的独立访客数」：详情接口先 ',
    'INSERT IGNORE',
    ' 写 ',
    't_view_record',
    '，写入成功（影响行数为 1）才执行 view_count = view_count + 1 原子自增（InnoDB 行锁保证并发准确）；'
    '重复访问被唯一索引挡下（影响行数为 0），不计数——访客反复刷新、浏览器前进后退、调试预览都不会把数字刷高。',
]):
    r.text = t

# P220 去重取舍
set_para_text(d.paragraphs[220],
              '• 去重未采用 Redis 方案：Redis 重启或清库后去重窗口会整体失效，还需额外定义'
              '“Redis 不可用时是否继续计数”的降级分支；本项目规模下由数据库唯一索引承担去重更简单、结果可复现且便于核对。'
              '同理也未引入定时批量落库——直接更新不存在缓存与库的最终一致性问题，也不会出现“刷新后数字不立即变化”的观感。')

# P222 预留扩展表收尾
set_para_text(d.paragraphs[222],
              '除上述 15 张业务表外，设计阶段还规划了敏感词表（sensitive_word）、文章版本历史表（article_history）'
              '与后台操作日志表（operation_log）三张预留扩展表，其 DDL 设计保留在数据库设计文档中；'
              '本期系统以人工审核、访问埋点等方式承担了对应职责，故未建表。')

# P402 结论章「不足」第四条
set_para_text(d.paragraphs[402],
              '受时间与条件限制，系统仍有以下不足：一是检索使用数据库 LIKE 匹配，数据量增大后性能会下降，'
              '方案中规划的 ngram 全文索引尚未启用；二是未实现敏感词过滤与短信、邮箱验证码，风控目前依赖后台人工审核；'
              '三是后台仪表盘以数字卡片为主，尚未引入趋势图表；'
              '四是文章浏览量虽已按独立访客去重，但仍为直接落库自增，未做 Redis 缓冲批量落库，高并发场景下写压力偏大。'
              '上述问题在文档中均给出了设计思路，属于后续迭代的优化方向而非功能缺陷。')

# ============ 2. article 表 view_count 行说明 ============
t_article = d.tables[3]
hit = 0
for row in t_article.rows:
    if row.cells[0].text.strip() == 'view_count':
        set_cell_text(row.cells[6], '浏览量（口径：读过该文的独立访客数）')
        hit += 1
assert hit == 1, f'view_count 行匹配 {hit} 次'

# ============ 3. 新增（15）view_record 小节 ============
p_num = d.paragraphs[211]     # （14）music 数据表（背景音乐表）：
p_cap = d.paragraphs[212]     # 表3.2.14 music 数据表
p_blank = d.paragraphs[213]   # 空行
p_idx = d.paragraphs[214]     # 索引：...
p_note = d.paragraphs[215]    # 设计说明：
p_items = d.paragraphs[216]   # 第一条 bullet
anchor = d.paragraphs[222]._p

anchor = clone_after(anchor, p_num, '（15）view_record 数据表（文章浏览记录表）：')
anchor = clone_after(anchor, p_cap, '表3.2.15 view_record 数据表')

tbl_el = copy.deepcopy(d.tables[13]._tbl)
anchor.addnext(tbl_el)
tbl = Table(tbl_el, p_cap._parent)
while len(tbl.rows) > 5:
    tbl._tbl.remove(tbl.rows[-1]._tr)
rows = [
    ['字段名', '类型', '长度', '允许空', '键/约束', '默认值', '说明'],
    ['id', 'BIGINT', '20', '否', 'PK', '—', '浏览记录ID'],
    ['article_id', 'BIGINT', '20', '否', 'FK / 联合唯一', '—', '被浏览文章ID'],
    ['visitor_key', 'VARCHAR', '64', '否', '联合唯一', '—', '访客标识：登录用户 u:{userId}；游客 ip:{客户端IP}'],
    ['create_time', 'DATETIME', '—', '否', '—', 'CURRENT_TIMESTAMP', '首次浏览时间'],
]
for ri, values in enumerate(rows):
    for ci, v in enumerate(values):
        set_cell_text(tbl.rows[ri].cells[ci], v)

anchor = clone_after(tbl_el, p_blank, '')
anchor = clone_after(anchor, p_idx,
                     '索引：PRIMARY(id)、uk_article_visitor(article_id, visitor_key)（唯一）、idx_article(article_id)。')
anchor = clone_after(anchor, p_note, '设计说明：')
for text in [
    '• 本表是浏览量的计数依据：文章详情接口先执行 INSERT IGNORE 写入本表，写入成功才让文章表的 view_count 加 1；'
    '写入被唯一索引挡下说明该访客已读过，不再计数。',
    '• 去重由唯一索引承担，不依赖 Redis 或内存状态，与点赞记录表、收藏表是同一套做法，'
    '天然支持多实例部署与服务重启后状态不丢。',
    '• 访客标识的取法：已登录用户取 u:{userId}，换设备、换网络仍识别为同一人；'
    '未登录游客取 ip:{客户端IP}，同一网络出口下的多个访客会被合并计一次，这是无登录态站点的常规近似。',
    '• 与访问日志埋点表的区别：后者是站点级访问记录（拦截器对门户请求异步记录 IP、省份、路径与 UA），'
    '本表是文章级的去重台账，两者互不替代；本表不预置种子数据，浏览量从真实访问累计。',
]:
    anchor = clone_after(anchor, p_items, text)

d.save(str(P))
print('正文 docx 已同步：15 张表口径 + 浏览量去重 + 新增（15）view_record 小节（表3.2.15）')
