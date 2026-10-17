# -*- coding: utf-8 -*-
# 修正版：把 TOC 域段落插到「目  录」标题之后（p24 之后）
import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = r"D:\quanbudaima\bi_ye_she_ji\docs\03-正文\毕业设计正文-简柚博客网站系统的设计与实现.docx"
doc = Document(DOCX)
ps = doc.paragraphs

# 定位「目  录」
toc_title_idx = next(i for i, p in enumerate(ps)
                     if re.sub(r"\s+", "", p.text) == "目录" and not p.style.name.startswith("toc"))
print("「目录」标题在段", toc_title_idx)

# 清理上一次误插的残留空段（p25 是空段且无分页符，属残留）
for i in range(toc_title_idx + 1, toc_title_idx + 3):
    el = ps[i]._element
    text = re.sub(r"\s+", "", ps[i].text)
    has_br = el.find(".//" + qn("w:br")) is not None
    has_sect = el.find(".//" + qn("w:sectPr")) is not None
    if not text and not has_br and not has_sect:
        print("删除残留空段", i)
        el.getparent().remove(el)

doc.save(DOCX)

# 重新打开，插到「目录」后
doc = Document(DOCX)
ps = doc.paragraphs
toc_title_idx = next(i for i, p in enumerate(ps)
                     if re.sub(r"\s+", "", p.text) == "目录" and not p.style.name.startswith("toc"))
anchor = ps[toc_title_idx]._element


def rpr(r):
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), "宋体")
    rPr.append(rf)
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag)
        e.set(qn("w:val"), "24")
        rPr.append(e)
    r.append(rPr)


def mkrun(p, kind=None, text=None):
    r = OxmlElement("w:r")
    rpr(r)
    if kind == "fld":
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), text)
        r.append(fc)
    elif kind == "instr":
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        r.append(it)
    elif kind == "txt":
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
    p.append(r)


p = OxmlElement("w:p")
pPr = OxmlElement("w:pPr")
jc = OxmlElement("w:jc")
jc.set(qn("w:val"), "left")
pPr.append(jc)
p.append(pPr)

mkrun(p, "fld", "begin")
mkrun(p, "instr", ' TOC \\o "1-2" \\h \\z \\u ')
mkrun(p, "fld", "separate")
mkrun(p, "txt", "【Ctrl+A 后按 F9 生成目录】")
mkrun(p, "fld", "end")

anchor.addnext(p)
doc.save(DOCX)

# 校验
doc = Document(DOCX)
xml = doc.part.element.xml
print("TOC 域:", "TOC" in xml, "| fldChar:", xml.count("fldChar"), "| bookmarkStart:", xml.count("bookmarkStart"))
ps = doc.paragraphs
for i in range(20, 32):
    print(i, repr(ps[i].style.name), repr(ps[i].text[:50]))
